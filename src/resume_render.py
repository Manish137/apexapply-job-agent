"""
Resume Renderer for ApexApply Job Agent
=======================================
Turns a plain-text resume (the format used in data/base_resumes) into clean HTML, PDF and DOCX.
Text layout rules: line 1 = name, line 2 = headline, line 3 = contact; ALL-CAPS lines are section
headings; lines starting with "•" are bullets; indented lines continue the previous line;
"Label: value" lines in SKILLS become bold-labelled rows; "(YYYY to YYYY|Present)" lines are roles.
"""

import html
import re
from pathlib import Path
from typing import List, Dict, Any, Optional

CSS = """body{font-family:Helvetica,Arial,sans-serif;font-size:10.5pt;color:#111;margin:0.55in 0.65in;line-height:1.35}
h1{font-size:20pt;margin:0}.tag{font-weight:600;color:#1f4e79;margin:2px 0}.contact{color:#444;margin-bottom:6px}
h2{font-size:11pt;text-transform:uppercase;letter-spacing:.06em;color:#1f4e79;border-bottom:1px solid #1f4e79;margin:12px 0 4px;padding-bottom:2px}
p{margin:3px 0}ul{margin:2px 0 4px 18px;padding:0}li{margin:2px 0}.role{font-weight:700;margin-top:8px}.sub{font-style:italic;color:#444}"""

ROLE_PATTERN = re.compile(r"\(\d{4} (to|-|–) (Present|\d{4})\)")


def _is_heading(line: str) -> bool:
    s = line.strip()
    return bool(s) and s == s.upper() and re.search(r"[A-Z]{3}", s) is not None and not s.startswith("•") and ":" not in s


def parse_resume_text(text: str) -> Dict[str, Any]:
    lines = [l.rstrip() for l in text.splitlines()]
    while lines and not lines[0].strip():
        lines.pop(0)
    header = [l.strip() for l in lines[:3]]
    sections: List[Dict[str, Any]] = []
    cur = None
    for raw in lines[3:]:
        if _is_heading(raw) and not ROLE_PATTERN.search(raw):
            cur = {"heading": raw.strip(), "items": []}
            sections.append(cur)
            continue
        if cur is None:
            continue
        s = raw.strip()
        if not s or set(s) <= set("═=-─"):
            cur["items"].append("")
            continue
        prev = cur["items"][-1] if cur["items"] else ""
        continuation = raw.startswith("  ") and prev and not s.startswith("•")
        summary_flow = cur["heading"] == "PROFESSIONAL SUMMARY" and prev and not s.startswith("•")
        if continuation or summary_flow:
            cur["items"][-1] = prev + " " + s
        else:
            cur["items"].append(s)
    return {"name": header[0] if header else "", "headline": header[1] if len(header) > 1 else "",
            "contact": header[2] if len(header) > 2 else "", "sections": sections}


def render_html(parsed: Dict[str, Any]) -> str:
    e = html.escape
    out = [f"<h1>{e(parsed['name'])}</h1><div class=tag>{e(parsed['headline'])}</div><div class=contact>{e(parsed['contact'])}</div>"]
    for sec in parsed["sections"]:
        out.append(f"<h2>{e(sec['heading'].title())}</h2>")
        in_list = False
        for it in sec["items"]:
            if not it:
                continue
            if it.startswith("•"):
                if not in_list:
                    out.append("<ul>")
                    in_list = True
                out.append(f"<li>{e(it[1:].strip())}</li>")
                continue
            if in_list:
                out.append("</ul>")
                in_list = False
            m = re.match(r"^([A-Z][A-Za-z &/]+?):\s+(.*)$", it)
            if sec["heading"] == "SKILLS" and m:
                out.append(f"<p><b>{e(m.group(1))}:</b> {e(m.group(2))}</p>")
            elif ROLE_PATTERN.search(it):
                out.append(f"<p class=role>{e(it)}</p>")
            elif sec["heading"] == "PROFESSIONAL EXPERIENCE":
                out.append(f"<p class=sub>{e(it)}</p>")
            else:
                out.append(f"<p>{e(it)}</p>")
        if in_list:
            out.append("</ul>")
    return (f"<!doctype html><html><head><meta charset=utf-8><title>{e(parsed['name'])} Resume</title>"
            f"<style>{CSS}</style></head><body>{''.join(out)}</body></html>")


def render_docx(parsed: Dict[str, Any], path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.7)
        s.top_margin = s.bottom_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    r = doc.add_paragraph().add_run(parsed["name"])
    r.bold = True
    r.font.size = Pt(18)
    doc.add_paragraph().add_run(parsed["headline"]).bold = True
    doc.add_paragraph(parsed["contact"])
    for sec in parsed["sections"]:
        doc.add_heading(sec["heading"].title(), level=2)
        for it in sec["items"]:
            if not it:
                continue
            m = re.match(r"^([A-Z][A-Za-z &/]+?):\s+(.*)$", it)
            if it.startswith("•"):
                doc.add_paragraph(it[1:].strip(), style="List Bullet")
            elif sec["heading"] == "SKILLS" and m:
                p = doc.add_paragraph()
                p.add_run(m.group(1) + ": ").bold = True
                p.add_run(m.group(2))
            elif ROLE_PATTERN.search(it):
                doc.add_paragraph().add_run(it).bold = True
            elif sec["heading"] == "PROFESSIONAL EXPERIENCE":
                doc.add_paragraph().add_run(it).italic = True
            else:
                doc.add_paragraph(it)
    doc.save(str(path))


def html_to_pdf(html_path: Path, pdf_path: Path) -> bool:
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page()
            page.goto(Path(html_path).resolve().as_uri())
            page.pdf(path=str(pdf_path), format="Letter", print_background=True)
            browser.close()
        return True
    except Exception:
        return False


def render_resume(text: str, out_base: Path, formats=("html", "pdf", "docx")) -> Dict[str, Optional[str]]:
    """Writes out_base.html/.pdf/.docx; returns paths (None for any format that failed)."""
    parsed = parse_resume_text(text)
    out_base.parent.mkdir(parents=True, exist_ok=True)
    result: Dict[str, Optional[str]] = {"html": None, "pdf": None, "docx": None}
    html_path = out_base.with_suffix(".html")
    html_path.write_text(render_html(parsed), encoding="utf-8")
    result["html"] = str(html_path)
    if "pdf" in formats and html_to_pdf(html_path, out_base.with_suffix(".pdf")):
        result["pdf"] = str(out_base.with_suffix(".pdf"))
    if "docx" in formats:
        try:
            render_docx(parsed, out_base.with_suffix(".docx"))
            result["docx"] = str(out_base.with_suffix(".docx"))
        except Exception:
            pass
    return result
